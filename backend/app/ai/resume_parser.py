import os
import re
import json
import logging
from pathlib import Path
from typing import Dict, List, Any, Optional

from app.core.config import settings

logger = logging.getLogger(__name__)

# Curated role requirements for gap analysis
ROLE_REQUIREMENTS: Dict[str, Dict[str, List[str]]] = {
    "Full Stack Developer": {
        "core": ["JavaScript", "TypeScript", "React", "Node.js", "Python", "FastAPI", "PostgreSQL", "Git", "REST API", "HTML", "CSS"],
        "recommended": ["Docker", "Tailwind CSS", "Redis", "Next.js", "GraphQL", "CI/CD", "AWS"]
    },
    "Backend Developer": {
        "core": ["Python", "FastAPI", "PostgreSQL", "SQL", "REST API", "Git", "Docker", "Database Design"],
        "recommended": ["Redis", "Kubernetes", "Microservices", "Celery", "RabbitMQ", "AWS", "CI/CD", "Linux"]
    },
    "Frontend Developer": {
        "core": ["JavaScript", "TypeScript", "React", "HTML", "CSS", "Tailwind CSS", "Git", "REST API"],
        "recommended": ["Next.js", "Redux", "Vue", "Responsive Design", "Webpack", "Vite", "Figma", "Jest"]
    },
    "Mobile App Developer": {
        "core": ["Flutter", "Dart", "React Native", "JavaScript", "Mobile UI", "REST API", "Git"],
        "recommended": ["Android", "iOS", "Swift", "Kotlin", "Firebase", "App Store Deployment", "SQLite"]
    },
    "Data Scientist": {
        "core": ["Python", "Pandas", "NumPy", "SQL", "Machine Learning", "Scikit-Learn", "Data Visualization", "Statistics"],
        "recommended": ["Deep Learning", "PyTorch", "TensorFlow", "Matplotlib", "Seaborn", "Jupyter", "Feature Engineering"]
    },
    "Machine Learning Engineer": {
        "core": ["Python", "PyTorch", "TensorFlow", "Machine Learning", "Deep Learning", "Docker", "Git", "Model Evaluation"],
        "recommended": ["MLOps", "Hugging Face", "Kubernetes", "ONNX", "FastAPI", "CUDA", "NLP", "Computer Vision"]
    },
    "DevOps Engineer": {
        "core": ["Linux", "Docker", "Kubernetes", "CI/CD", "Git", "Bash", "AWS", "Networking"],
        "recommended": ["Terraform", "Ansible", "Prometheus", "Grafana", "GitHub Actions", "Python", "Security"]
    },
    "Cloud Engineer": {
        "core": ["AWS", "Azure", "Cloud Architecture", "Linux", "Docker", "IAM", "Networking", "Git"],
        "recommended": ["Terraform", "Kubernetes", "GCP", "Serverless", "Lambda", "S3", "Monitoring"]
    },
    "QA / Automation Engineer": {
        "core": ["Python", "Selenium", "Postman", "API Testing", "Manual Testing", "Test Automation", "Git", "Bug Tracking"],
        "recommended": ["Playwright", "Cypress", "pytest", "CI/CD", "Performance Testing", "Jira", "SQL"]
    },
    "Cybersecurity Analyst": {
        "core": ["Network Security", "Linux", "Information Security", "Threat Analysis", "Wireshark", "Vulnerability Assessment"],
        "recommended": ["Penetration Testing", "SIEM", "Cryptography", "Firewalls", "Incident Response", "Python", "SOC"]
    },
    "UI/UX Designer": {
        "core": ["Figma", "UI Design", "UX Design", "Wireframing", "Prototyping", "User Research"],
        "recommended": ["Adobe XD", "Design Systems", "Usability Testing", "HTML/CSS", "Mobile Design", "Interaction Design"]
    }
}

ALL_COMMON_SKILLS = [
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C#", "Go", "Rust", "PHP", "Ruby", "Swift", "Kotlin", "Dart", "R", "SQL", "HTML", "CSS", "Bash",
    "React", "Vue", "Angular", "Next.js", "Nuxt.js", "Node.js", "Express", "FastAPI", "Django", "Flask", "Spring Boot", "ASP.NET", "Laravel",
    "PostgreSQL", "MySQL", "MongoDB", "SQLite", "Redis", "Elasticsearch", "Cassandra", "DynamoDB", "Firebase",
    "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Linux", "Terraform", "Ansible", "CI/CD", "GitHub Actions", "Jenkins", "Git",
    "PyTorch", "TensorFlow", "Scikit-Learn", "Pandas", "NumPy", "Keras", "OpenCV", "Hugging Face", "LLMs", "NLP",
    "Selenium", "Playwright", "Cypress", "pytest", "Postman", "Jira", "Figma", "Tailwind CSS", "Bootstrap", "REST API", "GraphQL", "WebSockets"
]


class ResumeParser:
    """Parses resumes from PDF / DOCX into structured candidate data and runs gap analysis."""

    def extract_text(self, file_path: str, file_type: str = "pdf") -> str:
        """Extracts text from PDF or DOCX file."""
        abs_path = Path(file_path).resolve()
        if not abs_path.exists():
            raise FileNotFoundError(f"File not found: {abs_path}")

        ext = file_type.lower().replace(".", "")
        if ext == "pdf":
            return self._extract_text_pdf(abs_path)
        elif ext in ["docx", "doc"]:
            return self._extract_text_docx(abs_path)
        else:
            # Fallback to plain text
            try:
                with open(abs_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            except Exception as e:
                logger.error(f"Error reading plain text file {abs_path}: {e}")
                return ""

    def _extract_text_pdf(self, path: Path) -> str:
        text = ""
        # 1. Try pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages_text = [page.extract_text() or "" for page in pdf.pages]
                text = "\n".join(pages_text).strip()
            if text:
                return text
        except Exception as e:
            logger.warning(f"pdfplumber extraction failed for {path}: {e}")

        # 2. Try pypdf fallback
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(pages_text).strip()
            if text:
                return text
        except Exception as e:
            logger.warning(f"pypdf extraction failed for {path}: {e}")

        return text

    def _extract_text_docx(self, path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_texts.append(" | ".join(row_cells))
            return "\n".join(paragraphs + table_texts)
        except Exception as e:
            logger.error(f"docx extraction failed for {path}: {e}")
            return ""

    def parse(self, raw_text: str) -> Dict[str, Any]:
        """
        Parses resume text into structured fields.
        First tries Gemini API if configured; falls back to robust local NLP heuristics.
        """
        if not raw_text.strip():
            return {
                "extracted_education": [],
                "extracted_skills": [],
                "extracted_projects": [],
                "extracted_certifications": [],
                "extracted_experience": [],
                "weak_sections": ["Resume document contains no readable text"],
                "improvement_suggestions": ["Upload a text-based PDF or DOCX file."],
            }

        # Try Gemini API if key is set
        if settings.GEMINI_API_KEY and settings.GEMINI_API_KEY.strip():
            try:
                gemini_result = self._parse_with_gemini(raw_text)
                if gemini_result and isinstance(gemini_result, dict):
                    return gemini_result
            except Exception as e:
                logger.warning(f"Gemini API resume parsing encountered an error, falling back to heuristics: {e}")

        return self._parse_with_heuristics(raw_text)

    def _parse_with_gemini(self, text: str) -> Optional[Dict[str, Any]]:
        from google import genai

        client = genai.Client(api_key=settings.GEMINI_API_KEY)
        prompt = (
            "You are an expert HR and Technical Resume Evaluator. Analyze the following resume text and extract "
            "comprehensive structured information in JSON format with exact keys:\n"
            "- extracted_education: list of objects with degree, institution, year, grade\n"
            "- extracted_skills: list of string skill names (technical, frameworks, tools, soft skills)\n"
            "- extracted_projects: list of objects with name, description, tech_stack (list of strings)\n"
            "- extracted_certifications: list of string certification names\n"
            "- extracted_experience: list of objects with role, company, duration, description\n"
            "- weak_sections: list of specific shortcomings or areas needing improvement (e.g., missing metrics, brief descriptions)\n"
            "- improvement_suggestions: list of actionable recommendations to enhance the resume for job applications\n\n"
            f"Resume Text:\n{text[:15000]}"
        )

        response_schema = {
            "type": "OBJECT",
            "properties": {
                "extracted_education": {
                    "type": "ARRAY",
                    "items": {
                        "type": "OBJECT",
                        "properties": {
                            "degree": {"type": "STRING"},
                            "institution": {"type": "STRING"},
                            "year": {"type": "STRING"},
                            "grade": {"type": "STRING"}
                        }
                    }
                },
                "extracted_skills": {"type": "ARRAY", "items": {"type": "STRING"}},
                "extracted_projects": {
                    "type": "ARRAY",
                    "items": {
                        "type": "OBJECT",
                        "properties": {
                            "name": {"type": "STRING"},
                            "description": {"type": "STRING"},
                            "tech_stack": {"type": "ARRAY", "items": {"type": "STRING"}}
                        }
                    }
                },
                "extracted_certifications": {"type": "ARRAY", "items": {"type": "STRING"}},
                "extracted_experience": {
                    "type": "ARRAY",
                    "items": {
                        "type": "OBJECT",
                        "properties": {
                            "role": {"type": "STRING"},
                            "company": {"type": "STRING"},
                            "duration": {"type": "STRING"},
                            "description": {"type": "STRING"}
                        }
                    }
                },
                "weak_sections": {"type": "ARRAY", "items": {"type": "STRING"}},
                "improvement_suggestions": {"type": "ARRAY", "items": {"type": "STRING"}},
            },
            "required": [
                "extracted_education",
                "extracted_skills",
                "extracted_projects",
                "extracted_certifications",
                "extracted_experience",
                "weak_sections",
                "improvement_suggestions"
            ]
        }

        interaction = client.interactions.create(
            model="gemini-3.8-flash",
            input=prompt,
            response_format=[
                {
                    "type": "text",
                    "mime_type": "application/json",
                    "schema": response_schema
                }
            ]
        )

        output = interaction.output_text
        if output:
            data = json.loads(output)
            return data
        return None

    def _parse_with_heuristics(self, text: str) -> Dict[str, Any]:
        """Robust rule-based parser used when LLM is unavailable or for offline testing."""
        lines = [line.strip() for line in text.splitlines() if line.strip()]

        # 1. Extract Skills
        extracted_skills = []
        text_lower = text.lower()
        for skill in ALL_COMMON_SKILLS:
            # Word boundary regex check
            pattern = r'\b' + re.escape(skill.lower()) + r'\b'
            if re.search(pattern, text_lower):
                extracted_skills.append(skill)

        # 2. Extract Education
        education = []
        edu_keywords = ["bachelor", "master", "bs", "ms", "b.s", "m.s", "phd", "matric", "intermediate", "fsc", "ics", "degree", "university", "institute", "college", "school"]
        for line in lines:
            if any(k in line.lower() for k in edu_keywords) and len(line) < 120:
                # Try to extract year
                year_match = re.search(r'\b(19\d{2}|20\d{2})\b', line)
                year = year_match.group(0) if year_match else None
                # Try to extract grade/CGPA
                gpa_match = re.search(r'(\d\.\d{1,2}(?:/\d\.\d{1,2})?|CGPA[:\s]*[\d\.]+|\b\d{2,3}%\b)', line, re.I)
                grade = gpa_match.group(0) if gpa_match else None

                education.append({
                    "degree": line,
                    "institution": "Academic Institution",
                    "year": year or "N/A",
                    "grade": grade or "N/A"
                })
        education = education[:5]  # limit to top matches

        # 3. Extract Experience
        experience = []
        exp_roles = ["intern", "engineer", "developer", "designer", "architect", "lead", "manager", "specialist", "assistant", "consultant"]
        for i, line in enumerate(lines):
            line_lower = line.lower()
            if any(r in line_lower for r in exp_roles) and len(line) < 90:
                # Potential next line might be description or company
                desc = lines[i + 1] if i + 1 < len(lines) else ""
                duration_match = re.search(r'((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s*\d{4}|\d{4})\s*[-–to]+\s*(Present|\d{4}|[a-z]+\s*\d{4})', line, re.I)
                duration = duration_match.group(0) if duration_match else "N/A"

                experience.append({
                    "role": line,
                    "company": "Company / Organization",
                    "duration": duration,
                    "description": desc
                })
        experience = experience[:5]

        # 4. Extract Projects
        projects = []
        proj_headings = ["project", "capstone", "portfolio", "application", "system", "app"]
        for i, line in enumerate(lines):
            if any(p in line.lower() for p in proj_headings) and len(line) < 80:
                proj_desc = lines[i + 1] if i + 1 < len(lines) else ""
                # Find project tech stack mentions
                p_skills = [s for s in extracted_skills if s.lower() in (line + " " + proj_desc).lower()]
                projects.append({
                    "name": line,
                    "description": proj_desc,
                    "tech_stack": p_skills
                })
        projects = projects[:5]

        # 5. Extract Certifications
        certifications = []
        cert_kws = ["certified", "certification", "certificate", "coursera", "udemy", "aws certified", "google certified", "cisco"]
        for line in lines:
            if any(c in line.lower() for c in cert_kws) and len(line) < 100:
                certifications.append(line)
        certifications = certifications[:5]

        # 6. Weak sections & suggestions
        weak_sections = []
        suggestions = []

        if not experience:
            weak_sections.append("No professional work experience or internships detected")
            suggestions.append("Add relevant internship, freelance, or academic project experience.")
        else:
            has_metrics = bool(re.search(r'\b\d+%\b|\b\d+x\b|\$\d+|\b\d+\s*(users|clients|requests|ms)\b', text, re.I))
            if not has_metrics:
                weak_sections.append("Experience section lacks quantifiable results or metrics")
                suggestions.append("Incorporate measurable impact (e.g., 'improved performance by 25%', 'served 10k users').")

        if len(extracted_skills) < 5:
            weak_sections.append("Limited technical skills listed")
            suggestions.append("Highlight primary programming languages, modern frameworks, and DevOps tools explicitly.")

        if not projects:
            weak_sections.append("No distinct project portfolio or academic projects highlighted")
            suggestions.append("Showcase 2-3 prominent hands-on projects with GitHub links and tech stack callouts.")

        if not certifications:
            weak_sections.append("No industry certifications listed")
            suggestions.append("Consider pursuing cloud, security, or domain-specific certifications to stand out.")

        suggestions.append("Use standard bullet points with strong action verbs (Architected, Developed, Optimized).")

        return {
            "extracted_education": education,
            "extracted_skills": extracted_skills,
            "extracted_projects": projects,
            "extracted_certifications": certifications,
            "extracted_experience": experience,
            "weak_sections": weak_sections,
            "improvement_suggestions": suggestions,
        }

    def analyze_skills_gap(self, extracted_skills: List[str], target_role_name: str) -> Dict[str, Any]:
        """
        Compares extracted resume skills against role requirements to produce missing skills,
        match score, and recommendations.
        """
        role_data = ROLE_REQUIREMENTS.get(target_role_name)
        if not role_data:
            # Fallback to general software engineer requirements if role not explicitly matched
            role_data = ROLE_REQUIREMENTS["Full Stack Developer"]

        extracted_lower = {s.lower() for s in extracted_skills}
        core_skills = role_data["core"]
        recommended_skills = role_data["recommended"]

        missing_core = [s for s in core_skills if s.lower() not in extracted_lower]
        missing_recommended = [s for s in recommended_skills if s.lower() not in extracted_lower]
        matching_core = [s for s in core_skills if s.lower() in extracted_lower]
        matching_recommended = [s for s in recommended_skills if s.lower() in extracted_lower]

        # Calculate match percentage
        total_weight = (len(core_skills) * 2) + len(recommended_skills)
        matched_weight = (len(matching_core) * 2) + len(matching_recommended)
        match_percentage = round((matched_weight / max(total_weight, 1)) * 100, 1)

        missing_formatted = []
        for s in missing_core:
            missing_formatted.append({"skill": s, "importance": "high", "type": "core"})
        for s in missing_recommended:
            missing_formatted.append({"skill": s, "importance": "medium", "type": "recommended"})

        return {
            "target_role": target_role_name,
            "match_percentage": match_percentage,
            "matching_skills": matching_core + matching_recommended,
            "missing_skills": missing_formatted,
            "core_skills_missing_count": len(missing_core),
            "total_missing_count": len(missing_formatted),
        }

resume_parser = ResumeParser()
