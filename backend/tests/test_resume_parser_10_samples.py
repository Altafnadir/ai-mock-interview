import io
import os
import sys
import pytest
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.ai.resume_parser import resume_parser, ROLE_REQUIREMENTS

TEMP_DIR = Path("backend/tests/fixtures/sample_resumes")
TEMP_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf_resume(filepath: Path, name: str, role: str, education: str, skills: str, experience: str, projects: str, certs: str):
    doc = SimpleDocTemplate(str(filepath), pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"<b>{name}</b> - {role}", styles["Title"]),
        Spacer(1, 10),
        Paragraph("<b>EDUCATION</b>", styles["Heading2"]),
        Paragraph(education, styles["Normal"]),
        Spacer(1, 10),
        Paragraph("<b>TECHNICAL SKILLS</b>", styles["Heading2"]),
        Paragraph(skills, styles["Normal"]),
        Spacer(1, 10),
        Paragraph("<b>PROFESSIONAL EXPERIENCE</b>", styles["Heading2"]),
        Paragraph(experience, styles["Normal"]),
        Spacer(1, 10),
        Paragraph("<b>PROJECTS</b>", styles["Heading2"]),
        Paragraph(projects, styles["Normal"]),
        Spacer(1, 10),
        Paragraph("<b>CERTIFICATIONS</b>", styles["Heading2"]),
        Paragraph(certs, styles["Normal"]),
    ]
    doc.build(story)


def create_docx_resume(filepath: Path, name: str, role: str, education: str, skills: str, experience: str, projects: str, certs: str):
    doc = Document()
    doc.add_heading(f"{name} - {role}", level=0)
    
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph(education)
    
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(skills)
    
    doc.add_heading("PROFESSIONAL EXPERIENCE", level=1)
    doc.add_paragraph(experience)
    
    doc.add_heading("PROJECTS", level=1)
    doc.add_paragraph(projects)
    
    doc.add_heading("CERTIFICATIONS", level=1)
    doc.add_paragraph(certs)
    
    doc.save(str(filepath))


# Define 10 distinct resumes across diverse roles and formats
RESUME_DEFINITIONS = [
    {
        "id": "sample_1_fullstack",
        "format": "pdf",
        "name": "Ahmed Khan",
        "role": "Full Stack Developer",
        "education": "Bachelor of Science in Software Engineering, GIMS PMAS Arid Agriculture University, 2024, CGPA 3.8",
        "skills": "Python, React, TypeScript, FastAPI, PostgreSQL, Docker, Git, REST API, HTML, CSS",
        "experience": "Full Stack Engineer at CloudTech (2023 - 2024). Designed microservices and React dashboards.",
        "projects": "AI Mock Interview Platform - Built video interview prep tool with React and FastAPI.",
        "certs": "AWS Certified Developer - Associate",
        "expected_skills": ["Python", "React", "FastAPI", "PostgreSQL", "Docker"]
    },
    {
        "id": "sample_2_backend",
        "format": "docx",
        "name": "Fatima Zahra",
        "role": "Backend Developer",
        "education": "BS Computer Science, NUST Islamabad, 2023",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Docker, SQL, REST API, Database Design, Git",
        "experience": "Backend Engineer at FinTech Corp. Architected high-throughput payment APIs and Redis cache.",
        "projects": "Distributed Ledger Service - Scalable transaction processing backend with FastAPI.",
        "certs": "PostgreSQL Certified Professional",
        "expected_skills": ["Python", "FastAPI", "PostgreSQL", "Redis", "Docker"]
    },
    {
        "id": "sample_3_frontend",
        "format": "pdf",
        "name": "Bilal Sheikh",
        "role": "Frontend Developer",
        "education": "Bachelor of Software Engineering, FAST-NUCES Lahore, 2024",
        "skills": "JavaScript, TypeScript, React, Next.js, HTML, CSS, Tailwind CSS, Git, REST API, Figma",
        "experience": "Frontend Developer Intern at Pixel Studio (Jan 2024 - Present). Built responsive UI components.",
        "projects": "E-Commerce Web App - Developed shopping portal with Next.js and Tailwind CSS.",
        "certs": "Meta Certified Front-End Developer",
        "expected_skills": ["JavaScript", "TypeScript", "React", "Next.js", "Tailwind CSS"]
    },
    {
        "id": "sample_4_qa_automation",
        "format": "pdf",
        "name": "Ayesha Malik",
        "role": "QA / Automation Engineer",
        "education": "BS Information Technology, University of the Punjab, 2023",
        "skills": "Python, Selenium, Playwright, Postman, pytest, API Testing, Manual Testing, Git, Jira",
        "experience": "QA Automation Engineer at QualityFirst. Authored automated test suites with pytest and Selenium.",
        "projects": "E2E Testing Suite - Automated end-to-end user regression tests for banking portal.",
        "certs": "ISTQB Certified Tester Foundation Level",
        "expected_skills": ["Python", "Selenium", "Playwright", "Postman", "pytest"]
    },
    {
        "id": "sample_5_devops",
        "format": "docx",
        "name": "Usman Tariq",
        "role": "DevOps Engineer",
        "education": "BS Computer Engineering, COMSATS Islamabad, 2022",
        "skills": "Linux, Docker, Kubernetes, CI/CD, AWS, Terraform, Bash, Git, GitHub Actions",
        "experience": "DevOps Specialist at InfraCloud. Deployed Kubernetes clusters and configured CI/CD pipelines.",
        "projects": "GitOps Infrastructure - Automated multi-region infrastructure provisioning using Terraform.",
        "certs": "Certified Kubernetes Administrator (CKA), AWS Solutions Architect",
        "expected_skills": ["Docker", "Kubernetes", "AWS", "Terraform", "Linux"]
    },
    {
        "id": "sample_6_mobile_developer",
        "format": "pdf",
        "name": "Zainab Ali",
        "role": "Mobile App Developer",
        "education": "BS Software Engineering, GIMS PMAS Arid University, 2024",
        "skills": "Flutter, Dart, React Native, JavaScript, Mobile UI, REST API, Git, Firebase",
        "experience": "Mobile Developer at AppVibe. Designed cross-platform mobile apps using Flutter and Firebase.",
        "projects": "Campus Navigator Mobile App - GPS navigation and timetable tracking in Flutter.",
        "certs": "Google Associate Android Developer",
        "expected_skills": ["Flutter", "Dart", "React Native", "Firebase", "Git"]
    },
    {
        "id": "sample_7_ml_ai_engineer",
        "format": "pdf",
        "name": "Hamza Raza",
        "role": "Machine Learning Engineer",
        "education": "MS Data Science, GIKI Topi, 2023 | BS Computer Science, 2021",
        "skills": "Python, PyTorch, TensorFlow, Machine Learning, Deep Learning, Docker, Git, Pandas, NumPy",
        "experience": "AI Research Associate at Cognitive Labs. Trained computer vision and LLM models.",
        "projects": "Facial Emotion Recognition - Deep learning model evaluating candidate stress and confidence.",
        "certs": "Deep Learning Specialization by DeepLearning.AI",
        "expected_skills": ["Python", "PyTorch", "TensorFlow", "Pandas", "NumPy"]
    },
    {
        "id": "sample_8_fresh_graduate",
        "format": "docx",
        "name": "Maryam Nawaz",
        "role": "Fresh Graduate (General HR)",
        "education": "Bachelor of Science in Software Engineering, GIMS Gujrat, 2024, CGPA 3.75",
        "skills": "Python, Java, C++, SQL, Git, HTML, CSS, Problem Solving, OOP, Database Design",
        "experience": "Academic Project Lead and Teaching Assistant at GIMS (2023 - 2024).",
        "projects": "Student Management Information System - Database driven academic portal.",
        "certs": "Python for Everybody Specialization",
        "expected_skills": ["Python", "Java", "C++", "SQL", "Git"]
    },
    {
        "id": "sample_9_cybersecurity",
        "format": "pdf",
        "name": "Danish Siddiqui",
        "role": "Cybersecurity Analyst",
        "education": "BS Cybersecurity, Air University Islamabad, 2023",
        "skills": "Network Security, Linux, Information Security, Wireshark, Vulnerability Assessment, Python, Cryptography",
        "experience": "Security Operations Center (SOC) Analyst at CyberShield. Analyzed intrusion alerts and network traffic.",
        "projects": "Vulnerability Scanner Tool - Python script automating port and vulnerability discovery.",
        "certs": "CompTIA Security+, CEH Certified Ethical Hacker",
        "expected_skills": ["Linux", "Python", "Network Security", "Wireshark"]
    },
    {
        "id": "sample_10_cloud_engineer",
        "format": "pdf",
        "name": "Sara Qureshi",
        "role": "Cloud Engineer",
        "education": "BS Computer Science, LUMS Lahore, 2023",
        "skills": "AWS, Azure, Docker, Kubernetes, Linux, IAM, Terraform, Git, Networking",
        "experience": "Cloud Solutions Associate at Enterprise Cloud. Configured AWS VPC, IAM policies, and S3 storage.",
        "projects": "Serverless Microservice Architecture - AWS Lambda and API Gateway deployment.",
        "certs": "AWS Certified Solutions Architect - Associate",
        "expected_skills": ["AWS", "Azure", "Docker", "Kubernetes", "Terraform"]
    }
]


@pytest.fixture(scope="module")
def sample_files():
    """Generates all 10 sample resumes (both PDF and DOCX) in a temporary fixture directory."""
    files = {}
    for r in RESUME_DEFINITIONS:
        ext = r["format"]
        path = TEMP_DIR / f"{r['id']}.{ext}"
        if ext == "pdf":
            create_pdf_resume(
                filepath=path,
                name=r["name"],
                role=r["role"],
                education=r["education"],
                skills=r["skills"],
                experience=r["experience"],
                projects=r["projects"],
                certs=r["certs"]
            )
        else:
            create_docx_resume(
                filepath=path,
                name=r["name"],
                role=r["role"],
                education=r["education"],
                skills=r["skills"],
                experience=r["experience"],
                projects=r["projects"],
                certs=r["certs"]
            )
        files[r["id"]] = path
    return files


@pytest.mark.parametrize("resume_spec", RESUME_DEFINITIONS)
def test_resume_parser_on_sample(resume_spec, sample_files):
    """
    Requirement S2: Verifies resume text extraction, structured parsing, 
    and skills gap analysis on 10 different sample resumes.
    """
    file_path = sample_files[resume_spec["id"]]
    assert file_path.exists(), f"Sample file {file_path} was not created."

    # 1. Text Extraction
    extracted_text = resume_parser.extract_text(str(file_path), resume_spec["format"])
    assert len(extracted_text) > 100, f"Extracted text too short for {resume_spec['id']}"
    assert resume_spec["name"].split()[0] in extracted_text, f"Candidate name not found in extracted text"

    # 2. Structured Parsing
    parsed_data = resume_parser.parse(extracted_text)
    assert isinstance(parsed_data, dict)
    
    # 3. Assert Skills Extracted
    extracted_skills = parsed_data.get("extracted_skills", [])
    assert len(extracted_skills) > 0, f"No skills extracted for {resume_spec['id']}"
    
    # Check that at least 2 of expected core skills were extracted
    overlap = set(s.lower() for s in extracted_skills).intersection(set(s.lower() for s in resume_spec["expected_skills"]))
    assert len(overlap) >= 2, f"Failed to extract key skills for {resume_spec['role']}. Extracted: {extracted_skills}"

    # 4. Assert Education Extracted
    education_entries = parsed_data.get("extracted_education", [])
    assert len(education_entries) > 0, f"No education entries found for {resume_spec['id']}"

    # 5. Assert Experience or Projects Extracted
    exp_entries = parsed_data.get("extracted_experience", [])
    proj_entries = parsed_data.get("extracted_projects", [])
    assert len(exp_entries) > 0 or len(proj_entries) > 0, f"Neither experience nor projects found for {resume_spec['id']}"

    # 6. Skills Gap Analysis for Target Role
    target_role = resume_spec["role"] if resume_spec["role"] in ROLE_REQUIREMENTS else "Full Stack Developer"
    gap_result = resume_parser.analyze_skills_gap(extracted_skills, target_role)
    assert "matching_skills" in gap_result
    assert "missing_skills" in gap_result
    assert "match_percentage" in gap_result

    assert isinstance(gap_result["match_percentage"], (int, float))
    assert gap_result["match_percentage"] >= 0.0
